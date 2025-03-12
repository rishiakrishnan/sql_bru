import os
import zipfile
import json
import xml.etree.ElementTree as ET
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook

def extract_metadata(filepath):
    # Placeholder function to extract metadata
    return {"metadata": "Sample metadata"}

def extract_docx_metadata(filepath):
    try:
        doc = Document(filepath)
        core_props = doc.core_properties
        metadata = {
            "author": core_props.author,
            "title": core_props.title,
            "subject": core_props.subject,
            "created": core_props.created,
            "modified": core_props.modified,
            "keywords": core_props.keywords,
        }
        return metadata
    except Exception as e:
        return {"error": str(e)}

def extract_pptx_metadata(filepath):
    try:
        prs = Presentation(filepath)
        core_props = prs.core_properties
        metadata = {
            "author": core_props.author,
            "title": core_props.title,
            "subject": core_props.subject,
            "created": core_props.created,
            "modified": core_props.modified,
            "keywords": core_props.keywords,
        }
        return metadata
    except Exception as e:
        return {"error": str(e)}

def extract_xlsx_metadata(filepath):
    try:
        wb = load_workbook(filepath)
        props = wb.properties
        metadata = {
            "author": props.creator,
            "title": props.title,
            "created": props.created,
            "modified": props.modified,
        }
        return metadata
    except Exception as e:
        return {"error": str(e)}

def extract_xml_scripts(filepath):
    try:
        tree = ET.parse(filepath)
        root = tree.getroot()
        scripts = []
        for elem in root.iter():
            if "script" in elem.tag.lower():
                scripts.append(elem.text)
        return scripts if scripts else "No scripts found"
    except Exception as e:
        return {"error": str(e)}

def scan_file(filepath):
    ext = filepath.split(".")[-1].lower()
    if ext == "docx":
        return extract_docx_metadata(filepath)
    elif ext == "pptx":
        return extract_pptx_metadata(filepath)
    elif ext == "xlsx":
        return extract_xlsx_metadata(filepath)
    elif ext in ["xml", "yaml"]:
        return extract_xml_scripts(filepath)
    else:
        return {"error": "Unsupported file type"}

if __name__ == "__main__":
    test_files = ["test.docx", "test.pptx", "test.xlsx", "test.xml"]
    for file in test_files:
        if os.path.exists(file):
            print(f"Metadata for {file}: {json.dumps(scan_file(file), indent=4)}")
        else:
            print(f"File {file} not found!")
